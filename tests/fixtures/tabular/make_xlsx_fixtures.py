# ruff: noqa: E501 (the XML of a workbook is written as it is stored)
"""Build the tabular adapter's XLSX fixtures: real small workbooks and hostile ones.

Run ``uv run python tests/fixtures/tabular/make_xlsx_fixtures.py`` to rewrite every ``.xlsx`` and
``.xlsm`` file under ``tests/fixtures/tabular/``. The writer is standard library only and the
parts are deterministic; the zip's deflate stream comes from zlib, whose bytes may differ between
zlib versions, so the tests compare the committed files to ``build()`` by their parts, not by bytes.

``--check`` validates the well-formed workbooks with the official reader (openpyxl) and nothing of
Neptune's::

    uv run --no-project --with openpyxl python tests/fixtures/tabular/make_xlsx_fixtures.py --check
"""

import datetime
import io
import random
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Final
from xml.sax.saxutils import escape

HERE: Final = Path(__file__).parent
MAIN: Final = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL: Final = "http://schemas.openxmlformats.org/package/2006/relationships"
DECL: Final = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
MEBIBYTE: Final = 1024 * 1024

# --- A small workbook writer ---------------------------------------------------------------------


def serial(when: datetime.datetime, *, epoch1904: bool = False) -> float:
    """The date serial of ``when`` in the 1900 system (the one with Excel's leap-year bug, so
    serials from March 1900 on are days since 1899-12-30) or the 1904 system."""
    base = datetime.datetime(1904, 1, 1) if epoch1904 else datetime.datetime(1899, 12, 30)
    delta = when - base
    return delta.days + delta.seconds / 86400


class Strings:
    """A shared strings table: every string once, in first-use order."""

    def __init__(self) -> None:
        self.items: list[str] = []
        self.raw: dict[int, str] = {}  # index -> a literal <si> body (rich text)

    def add(self, text: str) -> int:
        if text not in self.items:
            self.items.append(text)
        return self.items.index(text)

    def add_raw(self, plain: str, body: str) -> int:
        index = self.add(plain)
        self.raw[index] = body
        return index

    def xml(self) -> bytes:
        body = "".join(
            f"<si>{self.raw[i]}</si>"
            if i in self.raw
            else f'<si><t xml:space="preserve">{escape(text)}</t></si>'
            for i, text in enumerate(self.items)
        )
        return (
            f'{DECL}<sst xmlns="{MAIN}" count="{len(self.items)}"'
            f' uniqueCount="{len(self.items)}">{body}</sst>'
        ).encode()


def n(ref: str, value: object, style: int | None = None) -> str:
    s = f' s="{style}"' if style is not None else ""
    return f'<c r="{ref}"{s}><v>{value}</v></c>'


def s(ref: str, strings: Strings, text: str) -> str:
    return f'<c r="{ref}" t="s"><v>{strings.add(text)}</v></c>'


def i(ref: str, text: str) -> str:
    keep = ' xml:space="preserve"' if text != text.strip() else ""
    return f'<c r="{ref}" t="inlineStr"><is><t{keep}>{escape(text)}</t></is></c>'


def b(ref: str, flag: bool) -> str:
    return f'<c r="{ref}" t="b"><v>{int(flag)}</v></c>'


def e(ref: str, code: str) -> str:
    return f'<c r="{ref}" t="e"><v>{code}</v></c>'


def row(number: int | None, *cells: str) -> str:
    r = f' r="{number}"' if number is not None else ""
    return f"<row{r}>{''.join(cells)}</row>"


def worksheet(rows: list[str], before: str = "", after: str = "") -> bytes:
    return (
        f'{DECL}<worksheet xmlns="{MAIN}" xmlns:r="{REL}">{before}'
        f"<sheetData>{''.join(rows)}</sheetData>{after}</worksheet>"
    ).encode()


STYLES: Final = (
    f'{DECL}<styleSheet xmlns="{MAIN}">'
    '<numFmts count="1"><numFmt numFmtId="164" formatCode="yyyy\\-mm\\-dd\\ hh:mm"/></numFmts>'
    '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
    '<fills count="1"><fill><patternFill patternType="none"/></fill></fills>'
    '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
    '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    '<cellXfs count="4">'
    '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    '<xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    '<xf numFmtId="164" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    '<xf numFmtId="2" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>'
    "</cellXfs></styleSheet>"
).encode()


def workbook_xml(
    sheets: list[tuple[str, str]], *, epoch1904: bool = False, hidden: tuple[str, ...] = ()
) -> bytes:
    """``sheets`` are (name, relationship id)."""
    pr = '<workbookPr date1904="1"/>' if epoch1904 else '<workbookPr defaultThemeVersion="124226"/>'
    listed = "".join(
        f'<sheet name="{escape(name)}" sheetId="{k + 1}"'
        f'{" state=" + chr(34) + "hidden" + chr(34) if name in hidden else ""} r:id="{rid}"/>'
        for k, (name, rid) in enumerate(sheets)
    )
    return f'{DECL}<workbook xmlns="{MAIN}" xmlns:r="{REL}">{pr}<sheets>{listed}</sheets></workbook>'.encode()


def package(
    sheets: list[tuple[str, bytes]],
    *,
    strings: Strings | None = None,
    styles: bool = True,
    epoch1904: bool = False,
    hidden: tuple[str, ...] = (),
    macro: bool = False,
    extra_rels: str = "",
    extra_parts: dict[str, bytes] | None = None,
) -> dict[str, bytes]:
    """The parts of a workbook, in the order Excel writes them."""
    parts: dict[str, bytes] = {}
    main = (
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
        if macro
        else ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml")
    )
    overrides = [f'<Override PartName="/xl/workbook.xml" ContentType="{main}"/>']
    for k in range(len(sheets)):
        overrides.append(
            f'<Override PartName="/xl/worksheets/sheet{k + 1}.xml"'
            ' ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        )
    if styles:
        overrides.append(
            '<Override PartName="/xl/styles.xml"'
            ' ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
        )
    if strings is not None:
        overrides.append(
            '<Override PartName="/xl/sharedStrings.xml"'
            ' ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
        )
    parts["[Content_Types].xml"] = (
        f'{DECL}<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
        f"{''.join(overrides)}</Types>"
    ).encode()
    parts["_rels/.rels"] = (
        f'{DECL}<Relationships xmlns="{PKG_REL}"><Relationship Id="rId1"'
        f' Type="{REL}/officeDocument" Target="xl/workbook.xml"/></Relationships>'
    ).encode()
    rels = [
        f'<Relationship Id="rId{k + 1}" Type="{REL}/worksheet" Target="worksheets/sheet{k + 1}.xml"/>'
        for k in range(len(sheets))
    ]
    last = len(sheets)
    if styles:
        last += 1
        rels.append(f'<Relationship Id="rId{last}" Type="{REL}/styles" Target="styles.xml"/>')
    if strings is not None:
        last += 1
        rels.append(
            f'<Relationship Id="rId{last}" Type="{REL}/sharedStrings" Target="sharedStrings.xml"/>'
        )
    parts["xl/workbook.xml"] = workbook_xml(
        [(name, f"rId{k + 1}") for k, (name, _) in enumerate(sheets)],
        epoch1904=epoch1904,
        hidden=hidden,
    )
    parts["xl/_rels/workbook.xml.rels"] = (
        f'{DECL}<Relationships xmlns="{PKG_REL}">{"".join(rels)}{extra_rels}</Relationships>'
    ).encode()
    if styles:
        parts["xl/styles.xml"] = STYLES
    if strings is not None:
        parts["xl/sharedStrings.xml"] = strings.xml()
    for k, (_, xml) in enumerate(sheets):
        parts[f"xl/worksheets/sheet{k + 1}.xml"] = xml
    parts.update(extra_parts or {})
    return parts


def zipped(parts: dict[str, bytes]) -> bytes:
    """A zip of ``parts`` with fixed times, modes and order. Media are stored, as Excel stores
    already compressed pictures."""
    out = io.BytesIO()
    method = zipfile.ZIP_DEFLATED
    with zipfile.ZipFile(out, "w") as archive:
        for name, data in parts.items():
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED if name.startswith("xl/media/") else method
            info.create_system = 3
            info.external_attr = 0o600 << 16
            archive.writestr(info, data, compresslevel=9)
    return out.getvalue()


# --- The workbooks -------------------------------------------------------------------------------


def workorders_amr_fleet() -> dict[str, bytes]:
    """A CMMS work-order export for a warehouse mobile-robot fleet (Maximo-style columns)."""
    st = Strings()
    header = [
        "WONUM", "ASSETNUM", "DESCRIPTION", "STATUS", "WORKTYPE",
        "REPORTDATE", "LABORHRS", "SAFETY_CRITICAL", "FAILURECODE", "COMMENTS",
    ]  # fmt: skip
    cols = "ABCDEFGHIJ"
    rows = [row(1, *(s(f"{cols[k]}1", st, name) for k, name in enumerate(header)))]
    rich = st.add_raw(
        "Battery swap",
        '<r><t xml:space="preserve">Battery </t></r><r><rPr><b/></rPr><t>swap</t></r>'
        '<rPh sb="0" eb="7"><t>ignored</t></rPh>',
    )
    rows.append(
        row(
            2,
            s("A2", st, "WO-1001"), s("B2", st, "AMR-017"),
            s("C2", st, "Replace drive wheel bearing"), s("D2", st, "APPR"), s("E2", st, "CM"),
            n("F2", repr(serial(datetime.datetime(2026, 1, 8, 10, 0))), 2), n("G2", "3.5", 3),
            b("H2", True), s("I2", st, "BRG-WEAR"),
        )
    )  # fmt: skip
    rows.append(
        row(
            3,
            s("A3", st, "WO-1002"), s("B3", st, "AMR-004"), s("C3", st, "Lidar window cleaning"),
            s("D3", st, "COMP"), s("E3", st, "PM"), n("F3", "46030", 1), n("G3", "0.5", 3),
            b("H3", False), s("J3", st, "done on shift 2"),
        )
    )  # fmt: skip
    rows.append(
        row(
            4,
            s("A4", st, "WO-1003"), s("B4", st, "AMR-017"),
            f'<c r="C4" t="s"><v>{rich}</v></c>', s("D4", st, "WAPPR"), s("E4", st, "CM"),
            n("F4", "46031", 1), e("G4", "#N/A"), b("H4", True), s("I4", st, "BAT-LOW"),
        )
    )  # fmt: skip
    rows.append(
        row(
            5,
            s("A5", st, "WO-1004"), s("B5", st, "AMR-009"), s("C5", st, " spare part, lead 14 d "),
            s("D5", st, "CLOSE"), s("E5", st, "EM"), n("F5", "46032.5", 2),
            n("G5", "0.1"), '<c r="H5" s="1"/>', '<c r="I5" t="inlineStr"><is><t></t></is></c>',
            n("J5", "12345678901234567890123"),
        )
    )  # fmt: skip
    assets = [
        row(1, s("A1", st, "ASSETNUM"), s("B1", st, "MODEL"), s("C1", st, "SITE")),
        row(2, s("A2", st, "AMR-017"), s("B2", st, "tote-hauler-2"), s("C2", st, "DC-NORTH")),
        row(3, s("A3", st, "AMR-004"), s("B3", st, "tote-hauler-1"), s("C3", st, "DC-NORTH")),
    ]
    dimension = '<dimension ref="A1:J5"/><sheetViews><sheetView workbookViewId="0"/></sheetViews>'
    return package(
        [("Work Orders", worksheet(rows, dimension)), ("Assets", worksheet(assets))],
        strings=st,
        hidden=("Assets",),
    )


def changelog_manipulator_cell() -> dict[str, bytes]:
    """A change log of a robotic manipulator cell: inline strings, no shared strings or styles
    part, rows and cells with no reference, a gap, and parts the reader skips."""
    rows = [
        row(
            1,
            i("A1", "DATE"),
            i("B1", "CHANGE_ID"),
            i("C1", "CELL"),
            i("D1", "COMPONENT"),
            i("E1", "CHANGE"),
            i("F1", "APPROVED_BY"),
            i("G1", "REQUALIFIED"),
        ),
        row(
            2,
            n("A2", "46000"),
            i("B2", "CHG-0007"),
            i("C2", "cell-3"),
            i("D2", "gripper"),
            i("E2", "Swapped parallel gripper for vacuum end effector"),
            i("F2", "R. Okafor"),
            b("G2", True),
        ),
        # No row or cell references: positions are implied, A2's neighbours.
        "<row>"
        '<c t="n"><v>46002</v></c><c t="inlineStr"><is><t>CHG-0008</t></is></c>'
        '<c t="inlineStr"><is><t>cell-3</t></is></c><c t="inlineStr"><is><t>joint 4</t></is></c>'
        '<c t="inlineStr"><is><t>Recalibrated tool centre point</t></is></c></row>',
        # A gap: only A and F.
        row(4, n("A4", "46005"), i("F4", "M. Haddad")),
        row(5, n("A5", "46006"), i("B5", "CHG-0010"), i("C5", "cell-4"), b("G5", False)),
    ]
    before = (
        '<sheetPr><pageSetUpPr fitToPage="1"/></sheetPr><dimension ref="A1:G5"/>'
        '<cols><col min="1" max="1" width="12" customWidth="1"/></cols>'
    )
    after = '<mergeCells count="1"><mergeCell ref="D4:E4"/></mergeCells>'
    return package(
        [("Change Log", worksheet(rows, before, after))],
        styles=False,
        extra_parts={"xl/media/note.png": b"\x89PNG\r\n\x1a\n not read"},
    )


def epoch1904_quadruped() -> dict[str, bytes]:
    """A legged robot's inspection log in the 1904 date system (an old Mac Excel workbook)."""
    st = Strings()
    rows = [row(1, s("A1", st, "ASSET"), s("B1", st, "INSPECTED"), s("C1", st, "FINDING"))]
    when = [
        datetime.datetime(2026, 3, 2, 9, 30),
        datetime.datetime(2026, 3, 2, 14, 15),
        datetime.datetime(2026, 3, 9, 8, 0),
    ]
    names = ["Hip actuator FL", "Foot sensor RR", "Battery bay"]
    notes = ["OK", "WORN", "swollen cell"]
    for k, (name, stamp, note) in enumerate(zip(names, when, notes, strict=True)):
        r = k + 2
        rows.append(
            row(
                r, s(f"A{r}", st, name),
                n(f"B{r}", repr(serial(stamp, epoch1904=True)), 2), s(f"C{r}", st, note),
            )
        )  # fmt: skip
    return package([("Inspections", worksheet(rows))], strings=st, epoch1904=True)


def formulas_humanoid_energy() -> dict[str, bytes]:
    """A humanoid's energy budget: formulas with cached values, a shared and an array formula, a
    string and an error result, and a formula with no cached value."""
    st = Strings()
    rows = [
        row(1, s("A1", st, "joint"), s("B1", st, "kwh"), s("C1", st, "share")),
        row(2, s("A2", st, "hip"), n("B2", "1.5"), '<c r="C2"><f>B2/B5</f><v>0.375</v></c>'),
        row(
            3, s("A3", st, "knee"), n("B3", "2.5"),
            '<c r="C3"><f t="shared" ref="C3:C4" si="0">B3/B$5</f><v>0.625</v></c>',
        ),
        row(4, s("A4", st, "ankle"), n("B4", "0"), '<c r="C4"><f t="shared" si="0"/><v>0</v></c>'),
        row(
            5, s("A5", st, "total"), '<c r="B5"><f>SUM(B2:B4)</f><v>4</v></c>',
            '<c r="C5"><f>SUM(C2:C4)</f></c>',
        ),
        row(6, s("A6", st, "label"), '<c r="B6" t="str"><f>A2&amp;"-"&amp;A3</f><v>hip-knee</v></c>'),
        row(7, s("A7", st, "ratio"), '<c r="B7" t="e"><f>1/0</f><v>#DIV/0!</v></c>'),
        row(8, s("A8", st, "stacked"), '<c r="B8"><f t="array" ref="B8:B9">B2:B3*2</f><v>3</v></c>'),
        row(9, '<c r="B9"><v>5</v></c>'),
    ]  # fmt: skip
    return package([("Budget", worksheet(rows))], strings=st)


def truncated_workorders() -> bytes:
    data = zipped(workorders_amr_fleet())
    return data[: len(data) * 6 // 10]


def damaged_sheet_xml() -> dict[str, bytes]:
    """A valid zip whose sheet XML is cut inside its fourth row: the first rows must land."""
    st = Strings()
    rows = [
        row(1, s("A1", st, "id"), s("B1", st, "site")),
        row(2, n("A2", "1"), s("B2", st, "north")),
        row(3, n("A3", "2"), s("B3", st, "south")),
        row(4, n("A4", "3"), s("B4", st, "east")),
    ]
    xml = worksheet(rows)
    cut = xml.index(b'<c r="B4"') + 20
    return package([("Sites", xml[:cut])], strings=st)


def bomb_zeros_sheet() -> dict[str, bytes]:
    """A sheet padded with 150 MiB of spaces: about 1,000:1, over every size limit."""
    xml = worksheet([row(1, n("A1", "1"))], after=" " * (150 * MEBIBYTE))
    return package([("Pad", xml)])


def bomb_part_ratio() -> dict[str, bytes]:
    """Under the whole-file ratio but over the part's: 300 KiB of random bytes (a stored part)
    carry a sheet padded with 25 MiB of spaces."""
    xml = worksheet([row(1, n("A1", "1"))], after=" " * (25 * MEBIBYTE))
    noise = random.Random(7).randbytes(300 * 1024)
    return package([("Pad", xml)], extra_parts={"xl/media/noise.bin": noise})


def blowup_shared_strings() -> dict[str, bytes]:
    """3,000 shared strings and a sheet naming the first, a middle and the last of them: read
    with a lowered ``xlsx_max_shared_strings`` the later ones are not covered."""
    st = Strings()
    for k in range(3000):
        st.add(f"label-{k:04d}")
    rows = [
        row(
            1,
            *(f'<c r="{c}1" t="s"><v>{v}</v></c>' for c, v in (("A", 0), ("B", 1500), ("C", 2999))),
        )
    ]
    return package([("Labels", worksheet(rows))], strings=st)


def hostile_entities_sheet() -> dict[str, bytes]:
    """A sheet with a document type declaration and nested entities (a billion laughs)."""
    laughs = '<!ENTITY l0 "lol">' + "".join(
        f'<!ENTITY l{k} "{f"&l{k - 1};" * 10}">' for k in range(1, 10)
    )
    xml = (
        f'{DECL}<!DOCTYPE worksheet [{laughs}]><worksheet xmlns="{MAIN}"><sheetData>'
        '<row r="1"><c r="A1" t="inlineStr"><is><t>&l9;</t></is></c></row></sheetData></worksheet>'
    ).encode()
    return package([("Boom", xml)])


def macro_external_links() -> dict[str, bytes]:
    """A macro-enabled workbook that links to another workbook and an address."""
    st = Strings()
    rows = [
        row(1, s("A1", st, "tag"), s("B1", st, "value")),
        row(2, s("A2", st, "gain"), '<c r="B2"><f>[1]Sheet1!A1</f><v>2.5</v></c>'),
    ]
    extra = (
        f'<Relationship Id="rId90" Type="{REL}/externalLink" Target="externalLinks/externalLink1.xml"/>'
        f'<Relationship Id="rId91" Type="{REL}/hyperlink"'
        ' Target="https://example.invalid/never-followed" TargetMode="External"/>'
        f'<Relationship Id="rId92" Type="{REL}/vbaProject" Target="vbaProject.bin"/>'
    )
    link = (
        f'{DECL}<externalLink xmlns="{MAIN}"><externalBook xmlns:r="{REL}" r:id="rId1">'
        '<sheetNames><sheetName val="Sheet1"/></sheetNames></externalBook></externalLink>'
    ).encode()
    return package(
        [("Tuning", worksheet(rows))],
        strings=st,
        macro=True,
        extra_rels=extra,
        extra_parts={
            "xl/vbaProject.bin": b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(64),
            "xl/externalLinks/externalLink1.xml": link,
        },
    )


# fixture name -> the function that builds its parts (the bombs are large until zipped)
WORKBOOKS: Final[dict[str, Callable[[], dict[str, bytes]]]] = {
    "workorders_amr_fleet.xlsx": workorders_amr_fleet,
    "changelog_manipulator_cell.xlsx": changelog_manipulator_cell,
    "epoch1904_quadruped.xlsx": epoch1904_quadruped,
    "formulas_humanoid_energy.xlsx": formulas_humanoid_energy,
    "damaged_sheet_xml.xlsx": damaged_sheet_xml,
    "bomb_zeros_sheet.xlsx": bomb_zeros_sheet,
    "bomb_part_ratio.xlsx": bomb_part_ratio,
    "blowup_shared_strings.xlsx": blowup_shared_strings,
    "hostile_entities_sheet.xlsx": hostile_entities_sheet,
    "macro_external_links.xlsm": macro_external_links,
}


def build() -> dict[str, bytes]:
    files = {name: zipped(make()) for name, make in WORKBOOKS.items()}
    files["truncated_workorders.xlsx"] = truncated_workorders()
    return files


# --- Checking with the official reader -----------------------------------------------------------


def check() -> None:
    """Read the well-formed workbooks with openpyxl (and nothing of Neptune's)."""
    import openpyxl  # type: ignore[import-untyped,unused-ignore]

    def values(name: str, sheet: str, *, data_only: bool) -> list[list[object]]:
        book = openpyxl.load_workbook(HERE / name, data_only=data_only)
        return [[cell.value for cell in row] for row in book[sheet].iter_rows()]

    orders = values("workorders_amr_fleet.xlsx", "Work Orders", data_only=False)
    assert orders[0][:3] == ["WONUM", "ASSETNUM", "DESCRIPTION"], orders[0]
    assert orders[1][0] == "WO-1001" and orders[1][7] is True, orders[1]
    assert orders[1][5] == datetime.datetime(2026, 1, 8, 10, 0), orders[1][5]
    assert orders[3][2] == "Battery swap", orders[3][2]
    assert orders[3][6] == "#N/A" and orders[4][8] in (None, ""), orders[3:5]
    book = openpyxl.load_workbook(HERE / "workorders_amr_fleet.xlsx")
    assert book["Assets"].sheet_state == "hidden"

    log = values("changelog_manipulator_cell.xlsx", "Change Log", data_only=False)
    assert log[1][1] == "CHG-0007" and log[3][0] is not None and log[3][5] == "M. Haddad", log

    old = openpyxl.load_workbook(HERE / "epoch1904_quadruped.xlsx")
    assert old.epoch == openpyxl.utils.datetime.CALENDAR_MAC_1904, old.epoch
    inspected = [row[1].value for row in old["Inspections"].iter_rows(min_row=2)]
    assert inspected[0] == datetime.datetime(2026, 3, 2, 9, 30), inspected

    formulas = values("formulas_humanoid_energy.xlsx", "Budget", data_only=False)
    cached = values("formulas_humanoid_energy.xlsx", "Budget", data_only=True)
    assert formulas[1][2] == "=B2/B5" and cached[1][2] == 0.375, (formulas[1], cached[1])
    assert formulas[4][1] == "=SUM(B2:B4)" and cached[4][1] == 4, (formulas[4], cached[4])
    assert cached[4][2] is None, cached[4]
    assert cached[5][1] == "hip-knee" and cached[6][1] == "#DIV/0!", cached[5:7]
    sys.stdout.write("ok: openpyxl reads every well-formed workbook as the generator says\n")


def main() -> None:
    if "--check" in sys.argv:
        check()
        return
    for name, data in build().items():
        (HERE / name).write_bytes(data)


if __name__ == "__main__":
    main()
